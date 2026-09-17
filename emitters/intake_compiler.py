"""Compiles the unified Jarvis intake workbook (docs/templates/jarvis_intake_template.xlsx)
into Jarvis's existing contract shapes -- contracts/sources/<domain>/<source>.source.yaml,
contracts/models/<domain>.model.yaml, contracts/semantics/<domain>.gold.yaml -- plus one ODCS
v3 data contract per source. Replaces catalogue_compiler.py + silver_model_compiler.py +
gold_model_compiler.py's three-workbook flow with one workbook, validated up front against
contracts/schema/intake_spec.schema.json (ported from the "Data Platform Factory" reference
kit at D:/Projects/Data Platform - Jarvis/Design Docs).

Two-stage pipeline, same "preview then --write" discipline as the compilers it replaces:
  1. compile_workbook()      workbook -> canonical spec dict, schema-validated, cross-checked
  2. spec_to_contracts()     canonical spec -> Jarvis's three contract dicts + ODCS + warnings

KNOWN GAPS (translated as loudly-flagged warnings, never silently guessed -- CLAUDE.md rule 4):
  - Jarvis's silver engine (silver_transform.py) reads ONE bronze source per entity with a
    verbatim column SELECT -- it has no per-column SQL expression / multi-source UNION
    capability. A 08_Silver_Mappings row whose expression isn't an identity mapping, or an
    entity with mappings from more than one source, is accepted into the spec (so
    validate_spec and the human review at Freeze can see it) but is NOT executed by
    silver_transform.py -- flagged, not dropped, not silently "handled".
  - Jarvis's gold engine (gold_transform.py) builds SQL from structured metadata
    (source_fact + join_dimensions + grain + named metrics), not by executing a raw SELECT.
    09_Gold_Datasets' `sql` column is parsed only for its FROM-clause table name (source_fact);
    a JOIN anywhere in that SQL is flagged for manual join_dimensions_json authoring rather
    than guessed at.
  - check=natural_language DQ rules are never turned into SQL here -- they come back as
    Open Questions for the Freeze gate, same as every other agent in this project.
"""
from __future__ import annotations

import argparse
import json
import pathlib
import re
import sys
from typing import Any

import yaml
from jsonschema import Draft202012Validator
from openpyxl import load_workbook

REPO_ROOT = pathlib.Path(__file__).resolve().parent.parent
SCHEMA_PATH = REPO_ROOT / "contracts" / "schema" / "intake_spec.schema.json"
SOURCES_DIR = REPO_ROOT / "contracts" / "sources"
MODELS_DIR = REPO_ROOT / "contracts" / "models"
SEMANTICS_DIR = REPO_ROOT / "contracts" / "semantics"
ODCS_DIR = REPO_ROOT / "contracts" / "odcs"

YN = {"Y": True, "N": False, "YES": True, "NO": False, "TRUE": True, "FALSE": False}
BOOL_COLS = {"is_active", "has_header", "effective_date_in_filename", "is_mandatory",
             "nullable", "is_business_key", "pii"}
INT_COLS = {"data_source_id", "header_row_number", "ordinal", "precision", "scale"}
LIST_COLS = {"business_key", "merge_hash_columns", "dimensions", "recipients", "on"}


# --------------------------------------------------------------- stage 1: workbook -> spec

def _clean(col: str, val: Any) -> Any:
    if val is None or (isinstance(val, str) and val.strip() == ""):
        return None
    if col in BOOL_COLS:
        return YN.get(str(val).strip().upper())
    if col in INT_COLS:
        return int(val)
    if col in LIST_COLS:
        return [v.strip() for v in str(val).split(",") if v.strip()]
    if col == "params":
        out: dict[str, Any] = {}
        for part in str(val).split(";"):
            if "=" in part:
                k, v = part.split("=", 1)
                v = v.strip()
                out[k.strip()] = [x.strip() for x in v.split(",")] if "," in v and k.strip() in ("values", "allowed") else v
        return out
    if col == "scd_type":
        return int(val)
    if isinstance(val, str):
        return val.strip()
    return val


def _read_table(wb, sheet: str) -> list[dict[str, Any]]:
    ws = wb[sheet]
    headers = [c.value for c in ws[2]]
    rows = []
    for idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        rec: dict[str, Any] = {}
        for h, v in zip(headers, row):
            if h is None:
                continue
            cv = _clean(h, v)
            if cv is not None:
                rec[h] = cv
        rec["_row"] = idx
        rows.append(rec)
    return rows


def _set_path(d: dict, dotted: str, value: Any) -> None:
    keys = dotted.split(".")
    for k in keys[:-1]:
        d = d.setdefault(k, {})
    d[keys[-1]] = value


def compile_workbook(path: pathlib.Path) -> tuple[dict[str, Any], list[str]]:
    wb = load_workbook(path, data_only=True)
    errors: list[str] = []
    spec: dict[str, Any] = {"status": "draft"}

    for row in wb["01_Project"].iter_rows(min_row=3, values_only=True):
        key, val = row[0], row[1]
        if key and val not in (None, ""):
            _set_path(spec, key, str(val).strip())

    spec["data_systems"] = [{k: v for k, v in r.items() if k != "_row"} for r in _read_table(wb, "02_Data_Systems")]

    file_keys = {"file_name_pattern": "name_pattern", "file_name_regex": "name_regex", "file_format": "format",
                 "extension": "extension", "delimiter": "delimiter", "text_qualifier": "text_qualifier",
                 "has_header": "has_header", "header_row_number": "header_row_number", "row_tag": "row_tag",
                 "sheet_name": "sheet_name", "effective_date_in_filename": "effective_date_in_filename",
                 "effective_date_format": "effective_date_format"}
    load_keys = {"load_type", "frequency", "sla_cutoff_time", "is_mandatory", "watermark_column"}
    attrs = _read_table(wb, "04_Attributes")
    sources = []
    for r in _read_table(wb, "03_Data_Sources"):
        src: dict[str, Any] = {"file": {}, "load": {}, "attributes": []}
        for k, v in r.items():
            if k == "_row":
                continue
            if k in file_keys:
                src["file"][file_keys[k]] = v
            elif k in load_keys:
                src["load"][k] = v
            else:
                src[k] = v
        if "sla_cutoff_time" in src["load"] and not isinstance(src["load"]["sla_cutoff_time"], str):
            src["load"]["sla_cutoff_time"] = src["load"]["sla_cutoff_time"].strftime("%H:%M")
        src["attributes"] = sorted(
            [{k: v for k, v in a.items() if k not in ("_row", "data_source_code")}
             for a in attrs if a.get("data_source_code") == src.get("data_source_code")],
            key=lambda a: a.get("ordinal", 0))
        if not src["attributes"]:
            errors.append(f"03_Data_Sources row {r['_row']}: no attributes in 04_Attributes for '{src.get('data_source_code')}'")
        sources.append(src)
    spec["data_sources"] = sources

    source_codes = {s.get("data_source_code") for s in sources}
    system_codes = {s.get("data_system_code") for s in spec["data_systems"]}
    for a in attrs:
        if a.get("data_source_code") not in source_codes:
            errors.append(f"04_Attributes row {a['_row']}: unknown data_source_code '{a.get('data_source_code')}'")
    for s in sources:
        if s.get("data_system_code") not in system_codes:
            errors.append(f"03_Data_Sources: '{s.get('data_source_code')}' references unknown data_system_code '{s.get('data_system_code')}'")
    ids = [s.get("data_source_id") for s in sources]
    if len(ids) != len(set(ids)):
        errors.append("03_Data_Sources: data_source_id values must be unique")

    dq = _read_table(wb, "05_DQ_Rules")
    spec["dq_rules"] = [{k: v for k, v in r.items() if k != "_row"} for r in dq]
    for r in dq:
        if r.get("check") == "natural_language" and not r.get("description"):
            errors.append(f"05_DQ_Rules row {r['_row']}: natural_language rules need a description")

    ents = _read_table(wb, "06_Silver_Entities")
    cols = _read_table(wb, "07_Silver_Columns")
    entities = []
    for e in ents:
        ent = {k: v for k, v in e.items() if k != "_row"}
        ent["columns"] = [{k: v for k, v in c.items() if k not in ("_row", "entity")}
                          for c in cols if c.get("entity") == e.get("name")]
        if ent.get("scd_type") == 2 and not ent.get("merge_hash_columns"):
            errors.append(f"06_Silver_Entities row {e['_row']}: SCD2 entity '{e.get('name')}' needs merge_hash_columns")
        col_names = {c["name"] for c in ent["columns"]}
        for bk in ent.get("business_key", []):
            if bk not in col_names:
                errors.append(f"06_Silver_Entities row {e['_row']}: business key '{bk}' not in 07_Silver_Columns")
        entities.append(ent)
    maps = _read_table(wb, "08_Silver_Mappings")
    ent_cols = {e["name"]: {c["name"] for c in e["columns"]} for e in entities}
    for m in maps:
        if m.get("target_entity") not in ent_cols:
            errors.append(f"08_Silver_Mappings row {m['_row']}: unknown entity '{m.get('target_entity')}'")
        elif m.get("target_column") not in ent_cols[m["target_entity"]]:
            errors.append(f"08_Silver_Mappings row {m['_row']}: unknown column '{m.get('target_column')}'")
        if m.get("source") not in source_codes:
            errors.append(f"08_Silver_Mappings row {m['_row']}: unknown source '{m.get('source')}'")
    spec["silver_model"] = {"entities": entities,
                            "mappings": [{k: v for k, v in m.items() if k != "_row"} for m in maps]}

    datasets = _read_table(wb, "09_Gold_Datasets")
    ds_names = {d.get("name") for d in datasets}
    metrics = _read_table(wb, "10_Gold_Metrics")
    checks = _read_table(wb, "11_Business_Checks")
    for grp, label in ((metrics, "10_Gold_Metrics"), (checks, "11_Business_Checks")):
        for r in grp:
            if r.get("dataset") not in ds_names:
                errors.append(f"{label} row {r['_row']}: unknown dataset '{r.get('dataset')}'")
    spec["gold_model"] = {k: [{kk: vv for kk, vv in r.items() if kk != "_row"} for r in v]
                          for k, v in (("datasets", datasets), ("metrics", metrics), ("business_checks", checks))}

    spec.setdefault("governance", {})
    spec["governance"]["notifications"] = [{k: v for k, v in r.items() if k != "_row"}
                                           for r in _read_table(wb, "12_Notifications")]

    schema = json.loads(SCHEMA_PATH.read_text())
    for err in sorted(Draft202012Validator(schema).iter_errors(spec), key=lambda e: list(e.path)):
        loc = "/".join(str(p) for p in err.path) or "(root)"
        errors.append(f"schema: {loc}: {err.message}")

    ordered = {k: spec[k] for k in ["spec_version", "status", "project", "environment", "governance",
                                    "data_systems", "data_sources", "dq_rules", "silver_model", "gold_model"] if k in spec}
    return ordered, errors


# --------------------------------------------------------- stage 2: spec -> Jarvis contracts

_IMPLEMENTED_ATTR_CHECKS = {"not_null", "unique", "in_list", "range", "regex"}

# The unified schema's data_type enum (string/integer/bigint/decimal/double/boolean/date/
# timestamp/binary/json) doesn't fully match Jarvis's own contract vocabulary (bronze_loader.py's
# _DUCK_TYPE_MAP recognizes int/bigint/decimal/float/bool/timestamp, defaulting anything else to
# varchar) -- found by testing an integer column end-to-end: "integer" silently fell through to
# varchar, and a range DQ rule on it then failed at execute time comparing VARCHAR to DECIMAL.
# Normalize here, once, rather than editing bronze_loader.py's map to understand a second
# vocabulary.
_TYPE_MAP = {"integer": "int", "boolean": "bool", "double": "float", "binary": "string", "json": "string"}


def _jarvis_type(unified_type: str) -> str:
    return _TYPE_MAP.get(unified_type, unified_type)


def _source_contract(spec: dict, src: dict, warnings: list[str]) -> dict[str, Any]:
    domain = spec["project"]["domain"]
    client = spec["project"].get("client", "default")
    code = src["data_source_code"]
    attrs = sorted(src["attributes"], key=lambda a: a.get("ordinal", 0))

    file = src.get("file", {})
    ext = file.get("extension") or file.get("format", "csv")
    connection: dict[str, Any] = {
        "type": "file",
        # company first, kind of source second -- see emitters/landing.py
        "path": f"harness/landing/{domain}/structured/{code}/{code}_*.{ext}",
        "format": file.get("format", "csv"),
    }
    if file.get("delimiter"):
        connection["delimiter"] = file["delimiter"]
    connection["header"] = file.get("has_header", True)
    connection["encoding"] = file.get("encoding", "utf-8")

    contract: dict[str, Any] = {
        "$schema": "jarvis/source/v1",
        "source_id": code,
        "domain": domain,
        "client": client,
        "owner": src.get("owner") or spec["project"].get("owner", "data-eng@example.com"),
        "classification": "pii" if any(a.get("pii") for a in attrs) else "internal",
        "connection": connection,
    }

    load = src.get("load", {})
    arrival = {}
    if load.get("frequency"):
        arrival["cadence"] = load["frequency"]
    if load.get("sla_cutoff_time"):
        arrival["expected_by"] = load["sla_cutoff_time"]
    if arrival:
        contract["arrival"] = arrival

    # Jarvis's bronze layer is append-only by design (see bronze_loader.py's docstring) --
    # load_type from the workbook is informational, not yet a branch point in the engine.
    business_keys = [a["name"] for a in attrs if a.get("is_business_key")]
    contract["load"] = {"strategy": "append"}
    if business_keys:
        contract["load"]["primary_key"] = business_keys

    contract["schema"] = []
    for a in attrs:
        col: dict[str, Any] = {"name": a["name"], "type": _jarvis_type(a["data_type"]), "nullable": a.get("nullable", True)}
        if a.get("description"):
            col["description"] = a["description"]
        if a.get("pii"):
            col["pii"] = True
        if a.get("format"):
            col["format"] = a["format"]
        contract["schema"].append(col)

    file_rules = [r for r in spec.get("dq_rules", []) if r.get("target") == code and r.get("level") == "file"]
    min_rows = 100
    for r in file_rules:
        if r["check"] == "row_count_min":
            try:
                min_rows = int((r.get("params") or {}).get("min_rows", min_rows))
            except (TypeError, ValueError):
                pass
    unenforced_file_checks = [r["rule_id"] for r in file_rules if r["check"] != "row_count_min"]
    if unenforced_file_checks:
        warnings.append(
            f"{code}: file-level DQ rules {unenforced_file_checks} are recorded in dq_rules but not "
            "yet enforced by bronze_loader.py's FQC (only row_count_min maps to Jarvis's min_rows check today)."
        )
    contract["file_checks"] = {
        "min_rows": min_rows,
        "expected_columns": len(attrs),
        "row_count_deviation_pct": 40,
        "reject_on_schema_drift": True,
        "quarantine_path": f"harness/quarantine/{domain}/{code}/",
    }

    attr_rules = [r for r in spec.get("dq_rules", []) if r.get("target") == code and r.get("level") == "attribute"]
    quality_rules = []
    open_questions = []
    for r in attr_rules:
        check = r["check"]
        params = dict(r.get("params") or {})
        if check == "not_null":
            quality_rules.append({"rule": "not_null", "columns": [r["column"]], "severity": r["severity"]})
        elif check == "unique":
            quality_rules.append({"rule": "unique", "columns": [r["column"]], "severity": r["severity"]})
        elif check == "in_list":
            values = params.get("values", [])
            quality_rules.append({"rule": "accepted_values", "columns": [r["column"]], "severity": r["severity"],
                                   "params": {"values": values if isinstance(values, list) else [values]}})
        elif check == "range":
            p: dict[str, float] = {}
            if "min" in params:
                p["min"] = float(params["min"])
            if "max" in params:
                p["max"] = float(params["max"])
            quality_rules.append({"rule": "range", "columns": [r["column"]], "severity": r["severity"], "params": p})
        elif check == "regex":
            quality_rules.append({"rule": "regex", "columns": [r["column"]], "severity": r["severity"],
                                   "params": {"pattern": params.get("pattern", "")}})
        elif check == "natural_language":
            open_questions.append(f"{code}.{r.get('column', '(row)')}: \"{r.get('description', '')}\" "
                                   f"(DQ rule {r['rule_id']}) -- needs a human-approved concrete check before Freeze.")
        else:
            warnings.append(f"{code}: DQ rule {r['rule_id']} (check={check!r}) is not implemented by "
                             "bronze_loader.py's DQC yet -- recorded here for traceability, not enforced.")
    if quality_rules:
        contract["quality_rules"] = quality_rules
    if open_questions:
        contract["open_questions"] = open_questions

    return contract


def _entity_quality_rules(spec: dict, entity_name: str, warnings: list[str]) -> tuple[list[dict], list[str]]:
    """Silver-layer DQ rules: dq_rule rows with level=attribute, layer=silver, target=<entity>.
    Same rule vocabulary and translation as bronze's (see _source_contract) -- silver_transform.py
    gained an evaluator for this shape (see _eval_quality_rules there) specifically so these
    rules do something, not just get recorded."""
    rules = [r for r in spec.get("dq_rules", [])
             if r.get("level") == "attribute" and r.get("layer") == "silver" and r.get("target") == entity_name]
    quality_rules, open_questions = [], []
    for r in rules:
        check = r["check"]
        params = dict(r.get("params") or {})
        if check == "not_null":
            quality_rules.append({"rule": "not_null", "columns": [r["column"]], "severity": r["severity"]})
        elif check == "unique":
            quality_rules.append({"rule": "unique", "columns": [r["column"]], "severity": r["severity"]})
        elif check == "in_list":
            values = params.get("values", [])
            quality_rules.append({"rule": "accepted_values", "columns": [r["column"]], "severity": r["severity"],
                                   "params": {"values": values if isinstance(values, list) else [values]}})
        elif check == "range":
            p: dict[str, float] = {}
            if "min" in params:
                p["min"] = float(params["min"])
            if "max" in params:
                p["max"] = float(params["max"])
            quality_rules.append({"rule": "range", "columns": [r["column"]], "severity": r["severity"], "params": p})
        elif check == "regex":
            quality_rules.append({"rule": "regex", "columns": [r["column"]], "severity": r["severity"],
                                   "params": {"pattern": params.get("pattern", "")}})
        elif check == "natural_language":
            open_questions.append(f"{entity_name}.{r.get('column', '(row)')}: \"{r.get('description', '')}\" "
                                   f"(DQ rule {r['rule_id']}) -- needs a human-approved concrete check before Freeze.")
        else:
            warnings.append(f"{entity_name}: DQ rule {r['rule_id']} (check={check!r}) is not implemented by "
                             "silver_transform.py's quality-rule evaluator yet -- recorded, not enforced.")
    return quality_rules, open_questions


def _model_contract(spec: dict, warnings: list[str]) -> dict[str, Any]:
    entities = spec.get("silver_model", {}).get("entities", [])
    mappings = spec.get("silver_model", {}).get("mappings", [])
    dims, facts = [], []
    model_open_questions: list[str] = []

    for e in entities:
        by_source: dict[str, list[dict]] = {}
        for m in mappings:
            if m["target_entity"] == e["name"]:
                by_source.setdefault(m["source"], []).append(m)
        sources = list(by_source)
        if not sources:
            warnings.append(f"silver entity '{e['name']}' has no rows in 08_Silver_Mappings -- skipped entirely.")
            continue
        if len(sources) > 1:
            warnings.append(
                f"silver entity '{e['name']}' has mappings from {sources} -- Jarvis's silver_transform.py "
                f"only reads ONE bronze source per entity today (no multi-source conform/UNION yet). "
                f"Using '{sources[0]}'; rows mapped from {sources[1:]} are NOT loaded. This is the exact "
                "multi-provider conform gap flagged when this compiler was built -- see its module docstring."
            )
        source_table = sources[0]
        non_identity = [m["target_column"] for m in by_source[source_table]
                        if m["expression"].strip().lower() != m["target_column"].strip().lower()]
        if non_identity:
            warnings.append(
                f"silver entity '{e['name']}' columns {non_identity} have a non-identity expression in "
                "08_Silver_Mappings -- silver_transform.py does a verbatim column SELECT, so the bronze "
                "column name must already match the silver column name for these. The mapping expression "
                "itself is not applied."
            )

        quality_rules, entity_oqs = _entity_quality_rules(spec, e["name"], warnings)
        model_open_questions.extend(entity_oqs)

        if e["entity_type"] == "fact" or e.get("scd_type") == 0:
            fact: dict[str, Any] = {"name": e["name"], "grain": e["business_key"], "type": "transaction",
                                    "source": [source_table], "foreign_keys": [], "measures": []}
            if quality_rules:
                fact["quality_rules"] = quality_rules
            facts.append(fact)
        else:
            scd = e.get("scd_type", 1)
            dim: dict[str, Any] = {"name": e["name"], "type": "scd2" if scd == 2 else "scd1",
                                   "business_key": e["business_key"],
                                   "surrogate_key": e.get("surrogate_key") or f"{e['name']}_sk",
                                   "source": [source_table]}
            if scd == 2:
                # Fixed column names silver_transform.py's build_scd2_dimension() requires --
                # not user-configurable, same constants silver_model_compiler.py has always used.
                dim["effective_from"] = "row_start_date"
                dim["effective_to"] = "row_end_date"
                dim["current_flag"] = "row_is_current"
                if e.get("merge_hash_columns"):
                    dim["tracked_attributes"] = e["merge_hash_columns"]
                else:
                    warnings.append(f"SCD2 entity '{e['name']}' has no merge_hash_columns -- silver_transform.py "
                                     "requires this; the build will fail until it's set.")
            if quality_rules:
                dim["quality_rules"] = quality_rules
            dims.append(dim)

    model: dict[str, Any] = {
        "$schema": "jarvis/model/v1",
        "domain": spec["project"]["domain"],
        "client": spec["project"].get("client", "default"),
        "grain_statement": "One row per business entity or transaction, current-state conformed in silver.",
        "standardisation": {"date_format": "ISO-8601", "timezone": "UTC", "decimal_precision": [18, 4],
                            "string_case": "trim_only", "null_tokens": ["", "NULL", "N/A", "-"]},
        "dimensions": dims, "facts": facts,
        "conformance": {"shared_dimensions": [d["name"] for d in dims], "reject_orphan_fks": False},
    }
    if model_open_questions:
        model["open_questions"] = model_open_questions
    return model


_FROM_RE = re.compile(r"\bFROM\s+(?:silver\.)?(\w+)", re.IGNORECASE)
_JOIN_RE = re.compile(r"\bJOIN\b", re.IGNORECASE)


def _gold_contract(spec: dict, warnings: list[str]) -> dict[str, Any]:
    domain = spec["project"]["domain"]
    client = spec["project"].get("client", "default")
    gm = spec.get("gold_model", {})

    metrics = [{"name": m["name"], "expression": m["expression"], "grain": m.get("dimensions", []),
               "owner": m.get("owner")} for m in gm.get("metrics", [])]

    # check_type is carried through as executable fields (check_type/column/params/severity),
    # not just the "statement" every business_rules entry always had -- gold_transform.py gained
    # an evaluator (see _eval_business_checks) for the three check types the reference AWM slides
    # actually named (variance_vs_history, dimension_mapping_inconsistency,
    # new_or_missing_dimension); the other four in the schema's enum are recorded but flagged as
    # not executed, same "never silently guess" rule as everywhere else in this compiler.
    _IMPLEMENTED_BUSINESS_CHECKS = {"variance_vs_history", "dimension_mapping_inconsistency", "new_or_missing_dimension"}
    business_rules = []
    for bc in gm.get("business_checks", []):
        entry = {"id": bc["rule_id"], "statement": bc.get("description") or bc["check_type"],
                 "applies_to": [bc["dataset"]], "check_type": bc["check_type"], "dataset": bc["dataset"],
                 "column": bc.get("column"), "params": bc.get("params") or {}, "severity": bc.get("severity", "warning")}
        if bc["check_type"] not in _IMPLEMENTED_BUSINESS_CHECKS:
            warnings.append(f"business check {bc['rule_id']} (check_type={bc['check_type']!r}) is not "
                             "implemented by gold_transform.py's business-rule evaluator yet -- recorded, not enforced.")
        business_rules.append(entry)

    marts = []
    for ds in gm.get("datasets", []):
        sql = ds.get("sql", "")
        m = _FROM_RE.search(sql)
        if not m:
            warnings.append(f"gold dataset '{ds['name']}': couldn't find a FROM <table> in its sql -- "
                             "skipped, needs a manually-authored mart.")
            continue
        if _JOIN_RE.search(sql):
            warnings.append(f"gold dataset '{ds['name']}': its sql contains a JOIN -- gold_transform.py "
                             "needs an explicit join_dimensions_json, which this compiler does not derive "
                             "from raw SQL. Generated with join_dimensions=[] (no joins); edit "
                             f"contracts/semantics/{domain}.gold.yaml by hand to add the join(s) before Build.")
        ds_metrics = [mm["name"] for mm in gm.get("metrics", []) if mm["dataset"] == ds["name"]]
        marts.append({
            "name": ds["name"], "source_fact": m.group(1),
            "join_dimensions": [],  # see the JOIN warning above -- gold_transform.py reads this
                                     # key (already-parsed list), not a *_json string field
            "grain": [g.strip() for g in ds["grain"].split(",")] if isinstance(ds["grain"], str) else ds["grain"],
            "metrics": ds_metrics,
            "materialisation": ds.get("materialization", "table"),
        })

    return {
        "$schema": "jarvis/gold/v1", "domain": domain, "client": client,
        "consumers": ["bi_dashboard", "finance_export"],
        "metrics": metrics, "business_rules": business_rules, "marts": marts, "dashboards": [],
    }


def spec_to_contracts(spec: dict[str, Any]) -> dict[str, Any]:
    warnings: list[str] = []
    sources = [_source_contract(spec, s, warnings) for s in spec["data_sources"]]
    model = _model_contract(spec, warnings)
    gold = _gold_contract(spec, warnings)
    return {"sources": sources, "model": model, "gold": gold, "warnings": warnings}


ODCS_TYPES = {"string": "string", "integer": "integer", "bigint": "integer", "decimal": "number",
              "double": "number", "boolean": "boolean", "date": "date", "timestamp": "timestamp",
              "binary": "string", "json": "object"}


def to_odcs(spec: dict, src: dict) -> dict:
    """Minimal ODCS v3 contract for one source -- an industry-standard export alongside Jarvis's
    own YAML shape, not a replacement for it. Validate with `datacontract lint` if you adopt it
    for external sharing."""
    rules = [r for r in spec.get("dq_rules", []) if r.get("target") == src["data_source_code"]]
    props = []
    for a in src["attributes"]:
        p = {"name": a["name"], "logicalType": ODCS_TYPES[a["data_type"]], "physicalType": a["data_type"],
             "required": not a.get("nullable", True)}
        if a.get("is_business_key"):
            p["primaryKey"] = True
        if a.get("pii"):
            p["classification"] = "restricted"
        if a.get("description"):
            p["description"] = a["description"]
        col_rules = [r for r in rules if r.get("column") == a["name"]]
        if col_rules:
            p["quality"] = [{"type": "text", "name": r["rule_id"],
                             "description": r.get("description") or f"{r['check']} {r.get('params', '')}".strip()}
                            for r in col_rules]
        props.append(p)
    return {
        "apiVersion": "v3.0.2", "kind": "DataContract",
        "id": f"{spec['project']['code']}.{src['data_source_code']}",
        "name": src["entity_name"] + " - " + src.get("category", ""),
        "version": spec["spec_version"], "status": "draft", "domain": spec["project"]["domain"],
        "description": {"purpose": src.get("description", "")},
        "schema": [{"name": src["bronze_table"] if src.get("bronze_table") else src["data_source_code"],
                    "physicalName": src.get("bronze_table", src["data_source_code"]),
                    "logicalType": "object", "physicalType": "table", "properties": props}],
    }


def write_contracts(spec: dict[str, Any], compiled: dict[str, Any]) -> list[pathlib.Path]:
    domain = spec["project"]["domain"]
    written = []
    src_dir = SOURCES_DIR / domain
    src_dir.mkdir(parents=True, exist_ok=True)
    for contract, src in zip(compiled["sources"], spec["data_sources"]):
        path = src_dir / f"{contract['source_id']}.source.yaml"
        header = f"# Compiled from docs/templates/jarvis_intake_template.xlsx -- do not hand-edit.\n"
        path.write_text(header + yaml.safe_dump(contract, sort_keys=False, default_flow_style=False))
        written.append(path)

        odcs = to_odcs(spec, src)
        odcs_dir = ODCS_DIR / domain
        odcs_dir.mkdir(parents=True, exist_ok=True)
        odcs_path = odcs_dir / f"{src['data_source_code']}.odcs.yaml"
        odcs_path.write_text(yaml.safe_dump(odcs, sort_keys=False))
        written.append(odcs_path)

    MODELS_DIR.mkdir(parents=True, exist_ok=True)
    model_path = MODELS_DIR / f"{domain}.model.yaml"
    header = "# Contract schema: MODEL (silver). Compiled from docs/templates/jarvis_intake_template.xlsx -- do not hand-edit.\n"
    model_path.write_text(header + yaml.safe_dump(compiled["model"], sort_keys=False, default_flow_style=False))
    written.append(model_path)

    SEMANTICS_DIR.mkdir(parents=True, exist_ok=True)
    gold_path = SEMANTICS_DIR / f"{domain}.gold.yaml"
    header = "# Contract schema: GOLD / SEMANTIC. Compiled from docs/templates/jarvis_intake_template.xlsx -- do not hand-edit.\n"
    gold_path.write_text(header + yaml.safe_dump(compiled["gold"], sort_keys=False, default_flow_style=False))
    written.append(gold_path)

    return written


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default=str(REPO_ROOT / "docs" / "templates" / "jarvis_intake_template.xlsx"))
    ap.add_argument("--write", action="store_true")
    args = ap.parse_args()

    spec, errors = compile_workbook(pathlib.Path(args.workbook))
    if errors:
        print(f"FAILED with {len(errors)} error(s):")
        for e in errors:
            print("  -", e)
        sys.exit(1)

    compiled = spec_to_contracts(spec)
    print(f"OK: {len(compiled['sources'])} sources, {len(spec.get('dq_rules', []))} DQ rules, "
          f"{len(compiled['model']['dimensions'])} dims, {len(compiled['model']['facts'])} facts, "
          f"{len(compiled['gold']['marts'])} gold marts")

    domain = spec["project"]["domain"]
    domain_dir = SOURCES_DIR / domain
    if domain_dir.exists():
        existing = sorted(p.stem.removesuffix(".source") for p in domain_dir.glob("*.source.yaml"))
        new_ids = sorted(s["source_id"] for s in compiled["sources"])
        if existing and existing != new_ids:
            print(f"  ** WARNING ** domain {domain!r} already has contracts on disk: {existing}")
            print(f"  ** WARNING ** --write would replace matching files and LEAVE THE REST as-is "
                  f"(this compile only covers {new_ids}). If this domain already runs in production "
                  "(e.g. 'insurance'), double check project.domain in 01_Project before approving.")
    for w in compiled["warnings"]:
        print(f"  WARNING: {w}")
    for src in compiled["sources"]:
        for oq in src.get("open_questions", []):
            print(f"  OPEN QUESTION: {oq}")

    if args.write:
        written = write_contracts(spec, compiled)
        for p in written:
            print(f"  WROTE {p}")


if __name__ == "__main__":
    main()
