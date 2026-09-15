"""Compiles the Excel data catalogue (docs/templates/jarvis_data_catalogue_template.xlsx, or
any workbook matching its shape) into contracts/sources/<source_name>.source.yaml files.

This is the mechanism that lets a human fill in a spreadsheet instead of hand-writing YAML
(CLAUDE.md rule 1 still holds -- the compiler only ever emits what's IN the workbook; a missing
or inconsistent reference stops the run rather than guessing, rule 4).

Hierarchy read from the workbook, each level's rows carrying the parent's ID as an FK:
  Sources (source_id) -> Files (source_id FK) -> Tables (file_id FK)
  -> Columns (table_id FK), Quality Rules (table_id FK)

KNOWN LIMITATION, not silently papered over: the workbook does not yet capture unstructured-
source extras (extraction directives, bronze_schema, silver_shape) that
contracts/sources/policy_docs.source.yaml hand-authors beyond the generic source.schema.yaml
shape. Compiling an unstructured source produces a CONTRACT MISSING those sections. This script
refuses to overwrite an existing hand-authored contract when that would lose content -- it
writes new sources directly, but only ever previews a source that already has a contract on
disk, and says exactly what's missing.
"""
from __future__ import annotations

import argparse
import json
import pathlib
from typing import Any

import yaml
from openpyxl import load_workbook

REPO_ROOT = pathlib.Path(__file__).resolve().parent.parent
CONTRACTS_DIR = REPO_ROOT / "contracts" / "sources"


def _sheet_rows(ws) -> list[dict[str, Any]]:
    header = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:  # blank separator / end of data
            continue
        rows.append(dict(zip(header, row)))
    return rows


def _clean(v: Any) -> Any:
    """Excel gives us '' for a blank text cell and None for a truly empty one -- treat both
    as "not provided" so the compiler never emits an empty-string field into a contract."""
    if v is None or v == "":
        return None
    return v


def _split_list(v: Any) -> list[str] | None:
    v = _clean(v)
    if v is None:
        return None
    return [x.strip() for x in str(v).split(",") if x.strip()]


def load_catalogue(xlsx_path: pathlib.Path) -> dict[str, Any]:
    wb = load_workbook(xlsx_path, data_only=True)
    sources = _sheet_rows(wb["Sources"])
    files = _sheet_rows(wb["Files"])
    tables = _sheet_rows(wb["Tables"])
    columns = _sheet_rows(wb["Columns"])
    rules = _sheet_rows(wb["Quality Rules"])

    source_ids = {s["source_id"] for s in sources}
    file_ids = {f["file_id"] for f in files}
    table_ids = {t["table_id"] for t in tables}

    errors = []
    for f in files:
        if f["source_id"] not in source_ids:
            errors.append(f"Files row {f['file_id']!r} references unknown source_id {f['source_id']!r}")
    for t in tables:
        if t["file_id"] not in file_ids:
            errors.append(f"Tables row {t['table_id']!r} references unknown file_id {t['file_id']!r}")
    for c in columns:
        if c["table_id"] not in table_ids:
            errors.append(f"Columns row {c['column_name']!r} references unknown table_id {c['table_id']!r}")
    for r in rules:
        if r["table_id"] not in table_ids:
            errors.append(f"Quality Rules row references unknown table_id {r['table_id']!r}")
    if errors:
        raise ValueError(
            "Catalogue has broken references -- refusing to guess which row was meant:\n  "
            + "\n  ".join(errors)
        )

    return {"sources": sources, "files": files, "tables": tables, "columns": columns, "rules": rules}


def _build_connection(source: dict) -> dict[str, Any]:
    stype = source["source_type"]
    if stype == "file":
        conn = {"type": "file", "path": source["path"], "format": source["format"]}
        for k, excel_k in [("delimiter", "delimiter"), ("header", "header"), ("encoding", "encoding")]:
            if _clean(source.get(excel_k)) is not None:
                conn[k] = source[excel_k]
        return conn
    if stype == "database":
        conn = {"type": "database", "dialect": source["dialect"], "table": source["db_table"]}
        if _clean(source.get("incremental_column")) is not None:
            conn["incremental_column"] = source["incremental_column"]
        return conn
    if stype == "api":
        conn = {"type": "api", "endpoint": source["endpoint"]}
        if _clean(source.get("auth")) is not None:
            conn["auth"] = source["auth"]
        if _clean(source.get("secret_ref")) is not None:
            conn["secret_ref"] = source["secret_ref"]
        if _clean(source.get("pagination")) is not None:
            conn["pagination"] = source["pagination"]
        return conn
    if stype == "unstructured":
        conn = {"type": "unstructured", "path": source["path"]}
        formats = _split_list(source.get("unstructured_formats"))
        if formats:
            conn["formats"] = formats
        if _clean(source.get("max_file_mb")) is not None:
            conn["max_file_mb"] = source["max_file_mb"]
        return conn
    raise ValueError(f"Unknown source_type {stype!r} -- not one of file/database/api/unstructured")


def _build_file_checks(file_row: dict, is_unstructured: bool) -> dict[str, Any]:
    fc: dict[str, Any] = {}
    if is_unstructured:
        for k in ("min_files", "max_file_mb", "reject_zero_byte"):
            if _clean(file_row.get(k)) is not None:
                fc[k] = file_row[k]
    else:
        for k in ("min_rows", "max_rows", "expected_columns", "row_count_deviation_pct",
                   "size_deviation_pct", "reject_on_schema_drift"):
            if _clean(file_row.get(k)) is not None:
                fc[k] = file_row[k]
    if _clean(file_row.get("quarantine_path")) is not None:
        fc["quarantine_path"] = file_row["quarantine_path"]
    return fc


def compile_source(catalogue: dict, source_id: str) -> dict[str, Any]:
    source = next(s for s in catalogue["sources"] if s["source_id"] == source_id)
    files = [f for f in catalogue["files"] if f["source_id"] == source_id]
    if len(files) != 1:
        raise NotImplementedError(
            f"source {source_id} has {len(files)} Files rows -- multi-file/multi-table sources "
            "aren't implemented in the compiler yet (all four worked examples are 1:1:1)."
        )
    file_row = files[0]
    tables = [t for t in catalogue["tables"] if t["file_id"] == file_row["file_id"]]
    if len(tables) != 1:
        raise NotImplementedError(
            f"file {file_row['file_id']} has {len(tables)} Tables rows -- multi-table files "
            "aren't implemented in the compiler yet."
        )
    table = tables[0]
    columns = [c for c in catalogue["columns"] if c["table_id"] == table["table_id"]]
    rules = [r for r in catalogue["rules"] if r["table_id"] == table["table_id"]]
    is_unstructured = source["source_type"] == "unstructured"

    contract: dict[str, Any] = {
        "$schema": "jarvis/source/v1",
        "source_id": source["source_name"],
        "domain": source["domain"],
        "owner": source["owner"],
        "classification": source["classification"],
        "connection": _build_connection(source),
    }

    arrival = {}
    if _clean(source.get("arrival_cadence")) is not None:
        arrival["cadence"] = source["arrival_cadence"]
    for k in ("expected_by", "grace_minutes", "late_action"):
        if _clean(source.get(k)) is not None:
            arrival[k] = source[k]
    if arrival:
        contract["arrival"] = arrival

    load: dict[str, Any] = {}
    if _clean(table.get("load_strategy")) is not None:
        load["strategy"] = table["load_strategy"]
    pk = _split_list(table.get("primary_key"))
    if pk:
        load["primary_key"] = pk
    if _clean(table.get("watermark")) is not None:
        load["watermark"] = table["watermark"]
    dedupe = _split_list(table.get("dedupe_on"))
    if dedupe:
        load["dedupe_on"] = dedupe
    if load:
        contract["load"] = load

    if not is_unstructured:
        contract["schema"] = []
        for c in columns:
            col = {"name": c["column_name"], "type": c["type"], "nullable": bool(c["nullable"])}
            if _clean(c.get("description")):
                col["description"] = c["description"]
            if c.get("pii"):
                col["pii"] = True
            if _clean(c.get("format")):
                col["format"] = c["format"]
            contract["schema"].append(col)

    fc = _build_file_checks(file_row, is_unstructured)
    if fc:
        contract["file_checks"] = fc

    if rules:
        contract["quality_rules"] = []
        for r in rules:
            rule = {"rule": r["rule_type"], "columns": _split_list(r["columns"]),
                     "severity": r["severity"]}
            params = _clean(r.get("params_json"))
            if params:
                rule["params"] = json.loads(params)
            contract["quality_rules"].append(rule)

    sla = {}
    if _clean(source.get("sla_freshness_hours")) is not None:
        sla["freshness_hours"] = source["sla_freshness_hours"]
    if _clean(source.get("sla_completeness_pct")) is not None:
        sla["completeness_pct"] = source["sla_completeness_pct"]
    if sla:
        contract["sla"] = sla

    return contract


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", default=str(REPO_ROOT / "docs" / "templates" / "jarvis_data_catalogue_template.xlsx"))
    parser.add_argument("--write", action="store_true", help="Write contracts to disk (default: preview only)")
    args = parser.parse_args()

    catalogue = load_catalogue(pathlib.Path(args.workbook))
    for source in catalogue["sources"]:
        source_id = source["source_id"]
        source_name = source["source_name"]
        contract = compile_source(catalogue, source_id)
        yaml_text = yaml.safe_dump(contract, sort_keys=False, default_flow_style=False)

        existing_path = CONTRACTS_DIR / f"{source_name}.source.yaml"
        header = f"# {'-' * 70}\n# {source_id}  {source_name}  ({source['source_type']})\n# {'-' * 70}\n"

        if existing_path.exists():
            existing = existing_path.read_text()
            print(f"\n=== {source_name}: contract already exists at {existing_path} ===")
            if source["source_type"] == "unstructured":
                for missing in ("extraction:", "bronze_schema:", "silver_shape:"):
                    if missing in existing and missing not in yaml_text:
                        print(f"  NOT overwriting -- existing contract has {missing!r}, which the "
                              "workbook doesn't capture yet. Compiler output shown below for "
                              "comparison only.")
            print(header + yaml_text)
        else:
            print(f"\n=== {source_name}: new contract ===")
            print(header + yaml_text)
            if args.write:
                existing_path.write_text(header + yaml_text)
                print(f"  WROTE {existing_path}")


if __name__ == "__main__":
    main()
